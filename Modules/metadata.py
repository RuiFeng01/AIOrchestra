import os
from pathlib import Path

from PIL import Image, ExifTags
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import fitz  # PyMuPDF


def extract_pdf_metadata(filepath):
    try:
        doc = fitz.open(filepath)
        metadata = doc.metadata or {}
        return {
            "title": metadata.get("title"),
            "author": metadata.get("author"),
            "subject": metadata.get("subject"),
            "keywords": metadata.get("keywords"),
            "creator": metadata.get("creator"),
            "producer": metadata.get("producer"),
            "creation_date": metadata.get("creationDate"),
            "modification_date": metadata.get("modDate"),
            "trapped": metadata.get("trapped"),
            "format": metadata.get("format"),
            "encryption": metadata.get("encryption")
        }
    except Exception as e:
        return {"error": f"PDF metadata error: {e}"}


def extract_docx_metadata(filepath):
    try:
        doc = Document(filepath)
        props = doc.core_properties
        return {
            "author": props.author,
            "title": props.title,
            "subject": props.subject,
            "keywords": props.keywords,
            "comments": props.comments,
            "last_modified_by": props.last_modified_by,
            "created": str(props.created),
            "modified": str(props.modified),
            "category": props.category,
            "content_status": props.content_status,
            "identifier": props.identifier,
            "language": props.language,
            "version": props.version,
            "revision": props.revision
        }
    except Exception as e:
        return {"error": f"DOCX metadata error: {e}"}


def extract_xlsx_metadata(filepath):
    try:
        wb = load_workbook(filepath, read_only=True)
        props = wb.properties
        return {
            "author": props.creator,
            "title": props.title,
            "subject": props.subject,
            "keywords": props.keywords,
            "description": props.description,
            "last_modified_by": props.lastModifiedBy,
            "created": str(props.created),
            "modified": str(props.modified),
            "category": props.category,
            "content_status": props.contentStatus,
            "identifier": props.identifier,
            "language": props.language,
            "version": props.version,
            "revision": props.revision
        }
    except Exception as e:
        return {"error": f"XLSX metadata error: {e}"}


def extract_pptx_metadata(filepath):
    try:
        prs = Presentation(filepath)
        props = prs.core_properties
        return {
            "author": props.author,
            "title": props.title,
            "subject": props.subject,
            "keywords": props.keywords,
            "comments": props.comments,
            "last_modified_by": props.last_modified_by,
            "created": str(props.created),
            "modified": str(props.modified),
            "category": props.category,
            "content_status": props.content_status,
            "identifier": props.identifier,
            "language": props.language,
            "version": props.version,
            "revision": props.revision
        }
    except Exception as e:
        return {"error": f"PPTX metadata error: {e}"}


def extract_image_metadata(filepath):
    try:
        img = Image.open(filepath)
        exif_data = img._getexif() or {}
        readable_exif = {
            ExifTags.TAGS.get(tag, tag): value
            for tag, value in exif_data.items()
        }

        # Add general image info (format, size, mode, etc.)
        info = img.info if hasattr(img, "info") else {}
        return {
            "format": img.format,
            "mode": img.mode,
            "size": img.size,
            "dpi": info.get("dpi"),
            "software": info.get("Software") or readable_exif.get("Software"),
            "make": readable_exif.get("Make"),
            "model": readable_exif.get("Model"),
            "datetime_original": readable_exif.get("DateTimeOriginal"),
            "datetime_digitized": readable_exif.get("DateTimeDigitized"),
            "orientation": readable_exif.get("Orientation"),
            "fnumber": readable_exif.get("FNumber"),
            "exposure_time": readable_exif.get("ExposureTime"),
            "iso": readable_exif.get("ISOSpeedRatings"),
            "focal_length": readable_exif.get("FocalLength"),
            "white_balance": readable_exif.get("WhiteBalance"),
            "flash": readable_exif.get("Flash"),
            "gps_info": readable_exif.get("GPSInfo")
        }
    except Exception as e:
        return {"error": f"Image metadata error: {e}"}


def extract_file_metadata(filepath):
    ext = Path(filepath).suffix.lower()

    if ext == ".pdf":
        return extract_pdf_metadata(filepath)
    elif ext == ".docx":
        return extract_docx_metadata(filepath)
    elif ext == ".xlsx":
        return extract_xlsx_metadata(filepath)
    elif ext == ".pptx":
        return extract_pptx_metadata(filepath)
    elif ext in [".jpg", ".jpeg", ".tiff", ".png"]:
        return extract_image_metadata(filepath)
    else:
        return {"error": f"Unsupported file type or missing required library: {ext}"}


# Example usage
file_path = "metadata.pptx"  # Replace with your actual file path
metadata = extract_file_metadata(file_path)
for key, value in metadata.items():
    print(f"{key}: {value}")
